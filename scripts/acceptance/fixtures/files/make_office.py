"""生成附件验收用的 Office 文件和文本文件。每份文件里埋一个能核对的事实；产物提交在同一目录。
运行：python3 scripts/acceptance/fixtures/files/make_office.py（之后再跑 make-files.mts 生成 PDF）"""
import datetime, os, sys
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from PIL import Image, ImageDraw, ImageFont

out = os.path.dirname(os.path.abspath(__file__))

def docx():
    d = Document()
    d.add_heading("青柠项目周报", level=1)
    d.add_paragraph("项目代号：青柠 7 号。本周完成了支付页改版。")
    d.add_heading("预算", level=2)
    t = d.add_table(rows=3, cols=2)
    for r, (a, b) in enumerate([("项目", "金额"), ("设计", "38,000 元"), ("开发 | 测试", "128,000 元")]):
        t.cell(r, 0).text, t.cell(r, 1).text = a, b
    d.add_heading("下周计划", level=2)
    for item in ["联调退款接口", "灰度 5% 用户"]:
        d.add_paragraph(item, style="List Bullet")
    d.add_paragraph("负责人：林一舟")
    d.save(os.path.join(out, "weekly-report.docx"))

def xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["日期", "地区", "金额", "备注"])
    rows = [(datetime.date(2024, 3, 15), "华东", 4520.5, "首单"), (datetime.date(2024, 3, 16), "华南", 3180, ""), (datetime.date(2024, 3, 17), "华北", 0.1 + 0.2, "测试浮点")]
    for row in rows:
        ws.append(list(row))
    for c in ws["A"][1:]:
        c.number_format = "yyyy-mm-dd"
    ws["C6"] = "=SUM(C2:C4)"
    ws["B6"] = "合计"
    ws2 = wb.create_sheet("说明")
    ws2["A1"] = "数据口径：含税"
    ws3 = wb.create_sheet("内部")
    ws3["A1"] = "隐藏表里的暗号：北极星"
    ws3.sheet_state = "hidden"
    wb.save(os.path.join(out, "sales.xlsx"))

def pptx():
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "季度回顾"
    s.placeholders[1].text = "2024 Q1"
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "进展"
    s.placeholders[1].text_frame.text = "新增客户 42 家"
    s.placeholders[1].text_frame.add_paragraph().text = "续费率 91%"
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "风险"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    tbl.cell(0, 0).text, tbl.cell(0, 1).text = "风险", "应对"
    tbl.cell(1, 0).text, tbl.cell(1, 1).text = "供应商延期", "备选两家"
    s.notes_slide.notes_text_frame.text = "讲到这里提醒：合同 4 月 30 日到期"
    p.save(os.path.join(out, "review.pptx"))

def texts():
    with open(os.path.join(out, "cities-gbk.csv"), "wb") as f:
        f.write("城市,人口\n杭州,1252万\n苏州,1295万\n".encode("gb18030"))
    lines = [f"第 {i} 行：这是用来测试分段读取的填充文字，没有特别的信息。" for i in range(1, 3001)]
    lines[2400] = "第 2401 行：暗号是紫色长颈鹿。"
    with open(os.path.join(out, "long-notes.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    with open(os.path.join(out, "legacy.doc"), "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)

def scan_image():
    font = ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 54)
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    d.text((120, 200), "扫描件", font=font, fill="black")
    d.text((120, 320), "收据编号 A-1024", font=font, fill="black")
    d.text((120, 440), "金额：人民币 860 元", font=font, fill="black")
    img.save(os.path.join(out, "scan-page.png"))

docx(); xlsx(); pptx(); texts(); scan_image()
print("ok")
